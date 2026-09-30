import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu

TEMPLATE_PATH = r"C:\Users\apk05\.gemini\antigravity\brain\dedfdc13-12aa-458a-b065-3a9033f861d1\.user_uploaded\media_1790709480166.pptx"
OUTPUT_PATH = r"C:\Projects\QFlow\QFlow_Review_1_Presentation.pptx"
IMG_DIR = r"C:\Projects\QFlow\presentation_images"

def embed_images():
    prs = Presentation(TEMPLATE_PATH)

    # 1. Slide 5: Use Case Diagram
    slide5 = prs.slides[4]
    # Remove placeholder texts from slide 5 inside the left area
    shapes_to_remove = []
    for s in slide5.shapes:
        if s.has_text_frame and s.text in ['USE CASE DIAGRAM', 'Customer • Staff / Counter Operator • Administrator']:
            shapes_to_remove.append(s)
    for s in shapes_to_remove:
        sp = s._element
        sp.getparent().remove(sp)
    
    use_case_img = os.path.join(IMG_DIR, '01_use_case_diagram.png')
    slide5.shapes.add_picture(use_case_img, Emu(660000), Emu(1250000), width=Emu(7270000), height=Emu(4620000))

    # 2. Slide 6: ER Diagram
    slide6 = prs.slides[5]
    shapes_to_remove = []
    for s in slide6.shapes:
        if s.has_text_frame and s.text in ['ER DIAGRAM', 'User • Counter • Token • Service']:
            shapes_to_remove.append(s)
    for s in shapes_to_remove:
        sp = s._element
        sp.getparent().remove(sp)

    er_img = os.path.join(IMG_DIR, '02_er_diagram.png')
    slide6.shapes.add_picture(er_img, Emu(520000), Emu(1200000), width=Emu(7600000), height=Emu(4700000))

    # 3. Slide 15: Screenshots (8 slots)
    slide15 = prs.slides[14]
    # Remove placeholder texts
    placeholder_texts = [
        'STAFF LOGIN', 'Actual screen',
        'STAFF DASHBOARD', 'Counters + AI',
        'QR POSTER', 'Actual poster',
        'CUSTOMER FORM', 'Phone screen',
        'TOKEN STATUS', 'A101 + wait',
        'YOUR TURN', 'Live alert',
        'QUEUE / COUNTERS', 'Actions',
        'ANALYTICS', 'Charts',
        'Replace placeholders with the actual QFlow screenshots before submission.'
    ]
    shapes_to_remove = []
    for s in slide15.shapes:
        if s.has_text_frame and s.text.strip() in [t.strip() for t in placeholder_texts]:
            shapes_to_remove.append(s)
    for s in shapes_to_remove:
        sp = s._element
        sp.getparent().remove(sp)

    slots = [
        ('06_screen_staff_login.png', 457200, 1188720, 2743200, 2057400),
        ('07_screen_staff_dashboard.png', 3383280, 1188720, 2743200, 2057400),
        ('09_screen_qr_poster.png', 6309360, 1188720, 2743200, 2057400),
        ('03_screen_customer_form.png', 9235440, 1188720, 2468880, 2057400),
        ('04_screen_token_status.png', 457200, 3611880, 2743200, 2057400),
        ('05_screen_your_turn.png', 3383280, 3611880, 2743200, 2057400),
        ('08_screen_queue_counters.png', 6309360, 3611880, 2743200, 2057400),
        ('10_screen_analytics.png', 9235440, 3611880, 2468880, 2057400),
    ]

    for img_name, left, top, width, height in slots:
        img_path = os.path.join(IMG_DIR, img_name)
        if os.path.exists(img_path):
            slide15.shapes.add_picture(img_path, Emu(left + 20000), Emu(top + 20000), width=Emu(width - 40000), height=Emu(height - 40000))

    prs.save(OUTPUT_PATH)
    print(f"Presentation saved successfully to {OUTPUT_PATH}!")

if __name__ == '__main__':
    embed_images()
