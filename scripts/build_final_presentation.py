import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = r"C:\Users\apk05\.gemini\antigravity\brain\dedfdc13-12aa-458a-b065-3a9033f861d1\.user_uploaded\media_1790709480166.pptx"
FINAL_PPTX_PATH = r"C:\Projects\QFlow\QFlow_Review_1_Presentation.pptx"
IMG_DIR = r"C:\Projects\QFlow\presentation_images"

def build_presentation():
    prs = Presentation(TEMPLATE_PATH)

    # -------------------------------------------------------------------------
    # SLIDE 5: USE CASE DIAGRAM
    # -------------------------------------------------------------------------
    slide5 = prs.slides[4]
    shapes_to_remove_5 = []
    for s in slide5.shapes:
        if s.has_text_frame:
            txt = s.text.strip()
            if txt in ['USE CASE DIAGRAM', 'Customer • Staff / Counter Operator • Administrator']:
                shapes_to_remove_5.append(s)
    for s in shapes_to_remove_5:
        sp = s._element
        sp.getparent().remove(sp)

    use_case_img = os.path.join(IMG_DIR, '01_use_case_diagram.png')
    if os.path.exists(use_case_img):
        slide5.shapes.add_picture(use_case_img, Emu(660000), Emu(1250000), width=Emu(7270000), height=Emu(4620000))

    # -------------------------------------------------------------------------
    # SLIDE 6: ER DIAGRAM
    # -------------------------------------------------------------------------
    slide6 = prs.slides[5]
    shapes_to_remove_6 = []
    for s in slide6.shapes:
        if s.has_text_frame:
            txt = s.text.strip()
            if txt in ['ER DIAGRAM', 'MONGODB DATABASE DESIGN', 'User • Counter • Token • Service']:
                shapes_to_remove_6.append(s)
    for s in shapes_to_remove_6:
        sp = s._element
        sp.getparent().remove(sp)

    er_img = os.path.join(IMG_DIR, '02_er_diagram.png')
    if os.path.exists(er_img):
        slide6.shapes.add_picture(er_img, Emu(520000), Emu(1200000), width=Emu(7600000), height=Emu(4700000))

    # -------------------------------------------------------------------------
    # SLIDE 8: TECH STACK - Fix Blue MERN Text Contrast & Padding
    # -------------------------------------------------------------------------
    slide8 = prs.slides[7]
    for s in slide8.shapes:
        if s.has_text_frame:
            tf = s.text_frame
            tf.word_wrap = True
            for p in tf.paragraphs:
                if 'MERN =' in p.text:
                    p.font.color.rgb = RGBColor(79, 70, 229) # Clean Indigo
                    p.font.bold = True

    # -------------------------------------------------------------------------
    # SLIDE 15: SCREENSHOTS - Exactly 6 images in a clean 3x2 grid
    # -------------------------------------------------------------------------
    slide15 = prs.slides[14]
    
    # Remove all placeholder shapes (Rounded Rectangles 5, 8, 11, 14, 17, 20, 23, 26 and textboxes)
    shapes_to_remove_15 = []
    for s in slide15.shapes:
        if s.name not in ['Rectangle 1', 'TextBox 2', 'TextBox 3', 'TextBox 4']:
            shapes_to_remove_15.append(s)
    for s in shapes_to_remove_15:
        sp = s._element
        sp.getparent().remove(sp)

    # 6 Screenshots (3 columns x 2 rows)
    # Col widths: 3,550,000, Gap X: 220,000
    # Row heights: 2,200,000, Gap Y: 200,000
    col_x = [Emu(500000), Emu(4270000), Emu(8040000)]
    row_y = [Emu(1220000), Emu(3620000)]
    w = Emu(3550000)
    h = Emu(2200000)

    six_images = [
        ('03_screen_customer_form.png', 0, 0),
        ('04_screen_token_status.png', 1, 0),
        ('05_screen_your_turn.png', 2, 0),
        ('06_screen_staff_login.png', 0, 1),
        ('07_screen_staff_dashboard.png', 1, 1),
        ('08_screen_queue_counters.png', 2, 1),
    ]

    for img_file, c, r in six_images:
        img_path = os.path.join(IMG_DIR, img_file)
        if os.path.exists(img_path):
            slide15.shapes.add_picture(img_path, col_x[c], row_y[r], width=w, height=h)

    # -------------------------------------------------------------------------
    # SLIDE 17: REMAINING WORK & REFERENCES - Fix overlap
    # -------------------------------------------------------------------------
    slide17 = prs.slides[16]
    shapes_to_remove_17 = []
    for s in slide17.shapes:
        if s.has_text_frame and '17 slides total' in s.text:
            shapes_to_remove_17.append(s)
    for s in shapes_to_remove_17:
        sp = s._element
        sp.getparent().remove(sp)

    # Ensure all text boxes across all slides have clean padding
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                tf = shape.text_frame
                tf.word_wrap = True
                tf.margin_left = Emu(72000)
                tf.margin_right = Emu(72000)
                tf.margin_top = Emu(72000)
                tf.margin_bottom = Emu(72000)

    prs.save(FINAL_PPTX_PATH)
    print(f"Final presentation generated and saved to {FINAL_PPTX_PATH}")

if __name__ == '__main__':
    build_presentation()
