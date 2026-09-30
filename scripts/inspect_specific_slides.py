import os
from pptx import Presentation

TEMPLATE_PATH = r"C:\Users\apk05\.gemini\antigravity\brain\dedfdc13-12aa-458a-b065-3a9033f861d1\.user_uploaded\media_1790709480166.pptx"

prs = Presentation(TEMPLATE_PATH)

def inspect_slide(idx):
    slide = prs.slides[idx]
    print(f"\n==================== SLIDE {idx+1} ====================")
    for s in slide.shapes:
        text = s.text.replace('\n', ' ') if s.has_text_frame else '(no text)'
        text = text.encode('ascii', 'replace').decode('ascii')
        print(f"Shape: {s.name:20} | L:{s.left:8} T:{s.top:8} W:{s.width:8} H:{s.height:8} | {text[:70]}")

inspect_slide(7) # Slide 8 (Tech stack)
inspect_slide(11) # Slide 12 (Database integration)
inspect_slide(14) # Slide 15 (Screenshots)
inspect_slide(16) # Slide 17 (Remaining work & references)
