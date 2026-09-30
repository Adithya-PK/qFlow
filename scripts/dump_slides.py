import os
from pptx import Presentation

TEMPLATE_PATH = r"C:\Users\apk05\.gemini\antigravity\brain\dedfdc13-12aa-458a-b065-3a9033f861d1\.user_uploaded\media_1790709480166.pptx"

prs = Presentation(TEMPLATE_PATH)

for i, slide in enumerate(prs.slides):
    print(f"\n==================== SLIDE {i+1} ====================")
    for s in slide.shapes:
        text = s.text.replace('\n', ' ') if s.has_text_frame else '(no text)'
        has_table = s.has_table
        table_info = f"Table: {len(s.table.rows)}x{len(s.table.columns)}" if has_table else ""
        print(f"  Shape: {s.name:25} | L:{s.left:8} T:{s.top:8} W:{s.width:8} H:{s.height:8} | {table_info} | {text[:60]}")
